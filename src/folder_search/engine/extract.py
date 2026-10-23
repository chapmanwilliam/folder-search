"""Extract text from documents.

Returns a list of Page objects. For PDFs a Page is a real PDF page; for
other formats the whole document is a single Page numbered 0 (meaning "no
page number"). PDF pages with no text layer are counted, so a scanned file
can be OCR'd or reported rather than silently indexed as empty.
"""

from __future__ import annotations

import subprocess
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Page:
    number: int  # 1-based PDF page; 0 = not paginated
    text: str


@dataclass
class Extracted:
    pages: list[Page]
    empty_pages: int = 0  # PDF pages with no text layer (OCR candidates)

    @property
    def text_length(self) -> int:
        return sum(len(p.text) for p in self.pages)


class UnsupportedFormat(Exception):
    pass


SUPPORTED = {".pdf", ".docx", ".doc", ".rtf", ".md", ".txt", ".eml", ".msg"}


def extract(path: Path) -> Extracted:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".md", ".txt"):
        return _single(path.read_text(errors="replace"))
    if suffix in (".rtf", ".doc"):
        return _extract_textutil(path)
    if suffix == ".eml":
        return _extract_eml(path)
    if suffix == ".msg":
        return _extract_msg(path)
    raise UnsupportedFormat(suffix)


def _single(text: str) -> Extracted:
    text = text.strip()
    return Extracted([Page(0, text)] if text else [])


def _extract_pdf(path: Path) -> Extracted:
    import pymupdf

    pages: list[Page] = []
    empty = 0
    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise PermissionError("password-protected PDF")
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(Page(i, text))
            else:
                empty += 1
    return Extracted(pages, empty_pages=empty)


def _extract_docx(path: Path) -> Extracted:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _single("\n\n".join(parts))


def _extract_textutil(path: Path) -> Extracted:
    # textutil ships with macOS and reads RTF and legacy .doc well.
    out = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True,
        text=True,
        check=True,
    )
    return _single(out.stdout)


def _header_block(fields: list[tuple[str, str | None]]) -> str:
    return "\n".join(f"{k}: {v}" for k, v in fields if v)


def _html_to_text(html: str) -> str:
    import html as htmlmod
    import re

    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>", "\n", html)
    text = re.sub(r"<[^>]+>", " ", html)
    text = htmlmod.unescape(text)
    return re.sub(r"[ \t]+", " ", text)


def _extract_eml(path: Path) -> Extracted:
    from email import policy
    from email.parser import BytesParser

    with open(path, "rb") as f:
        msg = BytesParser(policy=policy.default).parse(f)
    header = _header_block(
        [
            ("From", msg["from"]),
            ("To", msg["to"]),
            ("Date", msg["date"]),
            ("Subject", msg["subject"]),
        ]
    )
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        content = body.get_content()
        text = _html_to_text(content) if body.get_content_subtype() == "html" else content
    attachments = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    if attachments:
        text += "\n\nAttachments: " + ", ".join(attachments)
    return _single(header + "\n\n" + text)


def _extract_msg(path: Path) -> Extracted:
    import extract_msg

    msg = extract_msg.openMsg(str(path))
    try:
        header = _header_block(
            [
                ("From", msg.sender),
                ("To", msg.to),
                ("Date", str(msg.date) if msg.date else None),
                ("Subject", msg.subject),
            ]
        )
        body = msg.body or ""
        if not body.strip() and msg.htmlBody:
            raw = msg.htmlBody
            body = _html_to_text(raw.decode("utf-8", "replace") if isinstance(raw, bytes) else raw)
        names = [a.longFilename or a.shortFilename for a in msg.attachments]
        names = [n for n in names if n]
        if names:
            body += "\n\nAttachments: " + ", ".join(names)
    finally:
        msg.close()
    return _single(header + "\n\n" + body)
