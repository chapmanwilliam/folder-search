from folder_search.engine.extract import extract

EML = b"""From: Sol <sol@example.com>
To: Counsel <c@example.com>
Subject: Instructions in Smith
Date: Mon, 1 Jun 2026 10:00:00 +0100
MIME-Version: 1.0
Content-Type: text/html; charset=utf-8

<p>Please advise on <b>limitation</b>.</p><script>x()</script>
"""


def test_eml_html_body_and_headers(tmp_path):
    p = tmp_path / "m.eml"
    p.write_bytes(EML)
    text = extract(p).pages[0].text
    assert "Subject: Instructions in Smith" in text
    assert "Please advise on" in text and "limitation" in text
    assert "<b>" not in text and "x()" not in text


def test_pdf_pages_and_empty_pages(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Page one text")
    doc.new_page()
    doc.new_page().insert_text((72, 72), "Page three text")
    p = tmp_path / "x.pdf"
    doc.save(p)
    ex = extract(p)
    assert [pg.number for pg in ex.pages] == [1, 3]
    assert ex.empty_pages == 1


def test_docx_includes_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Schedule of loss")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Care"
    t.rows[0].cells[1].text = "12,000"
    p = tmp_path / "s.docx"
    d.save(p)
    text = extract(p).pages[0].text
    assert "Schedule of loss" in text and "Care | 12,000" in text
